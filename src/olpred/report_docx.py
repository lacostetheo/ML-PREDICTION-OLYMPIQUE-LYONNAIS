"""Build the Word report in the format of the course group project (plain Calibri, numbered
questions, simple tables) and convert it to PDF.   Run:  python -m olpred.report_docx

Numbers that change with the data (top 3, European winner, previous results) are read from the
files produced by the pipeline, everything else comes from reports/tables/."""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPORTS = Path("reports")
FIG = REPORTS / "figures"
TAB = REPORTS / "tables"
OUT = REPORTS / "OL_ML_project.docx"
LEAGUES = {"E0": "Premier League", "SP1": "La Liga", "I1": "Serie A", "D1": "Bundesliga", "F1": "Ligue 1"}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_after = Pt(4)


def para(text="", bold=False, italic=False, size=11, align=None, after=4, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    if align:
        p.alignment = align
    # @@bold@@ segments inside the text
    parts = text.split("@@")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or i % 2 == 1
        r.italic = italic
        r.font.size = Pt(size)
    return p


def title(text):
    para(text, bold=True, size=16, after=2, keep=True)


def h1(text):
    para(text, bold=True, size=13, after=2, keep=True)


def h2(text):
    para(text, bold=True, size=11, after=2, keep=True)


def bullets(items):
    for it in items:
        para("- " + it, after=1)
    para("", after=2)


def formula(text, note=None):
    p = para(text, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p.paragraph_format.space_before = Pt(4)
    for r in p.runs:
        r.font.name = "Cambria Math"
        r.font.size = Pt(11.5)
    if note:
        para(note, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)


def _borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def table(rows, widths=None, size=9.5, bold_rows=(), align_right_from=1):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _borders(t)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(size)
            r.bold = i == 0 or i in bold_rows
            if widths:
                c.width = Cm(widths[j])
    if widths:
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
        for j, wcm in enumerate(widths):
            t.columns[j].width = Cm(wcm)
    if len(rows) <= 16:  # keep short tables on one page
        for row in t.rows[:-1]:
            for c in row.cells:
                for p_ in c.paragraphs:
                    p_.paragraph_format.keep_with_next = True
    para("", after=6)
    return t


def figure(name, width_cm=14.0, caption=None):
    doc.add_picture(str(FIG / name), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        para(caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def pct(x, d=0):
    return f"{100 * x:.{d}f}%"


# ------------------------------------------------------------------ data pulled from the pipeline
def site_data():
    t = Path("site/data.js").read_text()
    return json.loads(t[len("window.OLDATA = "):].rstrip().rstrip(";"))


D = site_data()
final = pd.read_csv(TAB / "final_test_table.csv", index_col=0)
ucl = pd.read_csv(TAB / "europe_ucl_probabilities.csv")
euro = pd.read_csv(TAB / "europe_walk_forward.csv", index_col=0)
euro_meta = json.loads((TAB / "europe_ucl_meta.json").read_text())
prev = D["prev"]
ol = D["leagues"]["F1"]["sim"]["summary"]["Lyon"]  # exp_pts,p10,p90,exp_pos,title,cl,top4,top6,europe,rel,pts_now,played
ol_pos = D["leagues"]["F1"]["sim"]["pos_probs"]["Lyon"]

# ================================================================== TITLE + DISCLAIMER
title("Machine Learning Project: Predicting the Victories of Olympique Lyonnais")
para("Théo Lacoste, EDHEC Business School  |  October 2026", size=10.5, after=10)

para("@@Disclaimer.@@ I hope the ML does not predict well, because what the f***: OL will be 5th this season. "
     f"(The simulation's average final position for OL is {ol[3]:.1f}. So much for my hopes.)", after=8)
para("Joke aside, let's start.", bold=True, after=10)

# ================================================================== INTRODUCTION
h1("Introduction")
para("This project asks a simple question: @@can a machine-learning model predict what Olympique Lyonnais (OL) will do, "
     "and is it better than simple references such as an Elo rating or the betting market?@@ "
     "Football is a good test. A match is decided by a handful of goals, the stronger team wins only about one time in two, "
     "and one game in four ends in a draw. The same question matters in finance: how much can a model extract from public data, "
     "how do we measure it reliably, and does it survive comparison with a liquid market?")
para("OL alone plays about 35 league matches a year, which is far too few to train a model. So I train on @@all five major leagues "
     "(Premier League, La Liga, Serie A, Bundesliga, Ligue 1; 29,239 matches from 2010/11 to October 2026)@@ and use OL as a case study.")
h2("The model I chose")
para("@@Multinomial logistic regression@@ on 100 input variables (Elo rating, squad market value, recent form, shots, schedule), "
     "re-fitted every season. For each match it gives three probabilities: home win, draw, away win. "
     "I also tested random forest, gradient boosting, a support-vector machine and neural networks (all in scikit-learn). "
     "None of them is significantly better, so I kept the simplest and most explainable one (section 4).")
h2("Three rules I followed")
bullets(["@@No leakage:@@ a feature used to predict a match on date T contains nothing dated T or later. An automated test checks it.",
         "@@Temporal validation:@@ train on the past, judge on the future, never on a random split.",
         "@@Honest evidence:@@ every claimed improvement comes with a confidence interval, and negative results are reported as negative."])
h2("The answer in five lines")
bullets([f"Team strength (Elo and squad value) explains almost everything that can be predicted.",
         f"The chosen model has a test log loss of {final.loc['LogReg (final)', 'log_loss']:.3f}, against {final.loc['Class frequency (naive)', 'log_loss']:.3f} for a model that ignores team quality.",
         f"The bookmaker is better ({final.loc['Bookmaker Bet365 (reference)', 'log_loss']:.3f}), and mixing my model into its prices adds nothing.",
         f"For OL in Ligue 1 2026/27: {ol[0]:.1f} expected points, average final position {ol[3]:.1f}, {pct(ol[6])} chance of the top 4, {pct(ol[4])} of the title.",
         "No clear favourite for the Champions League: five clubs sit around 13% each."])
page_break()

# ================================================================== 1
h1("1. Research question and targets")
para("@@Question:@@ can machine learning predict the sporting performance of OL from team-level, match-level and contextual data, "
     "and does it beat simple references (Elo, the betting market)?")
para("My first idea was to feed player data to a neural network to predict OL's final ranking. I dropped it: a final ranking is one observation "
     "per team and season (about 20 per league), and OL has only about 350 matches in ten years. So I predict the @@match result@@ instead, "
     "where there are about 29,000 examples, and I obtain season quantities by simulation.")
table([["Candidate target", "Decision"],
       ["Win / draw / loss probabilities of a match", "Main target (about 29,000 observations)"],
       ["Expected points", "Derived: 3 x P(win) + P(draw)"],
       ["Final points, rank, Europe, relegation", "Not trained (too few observations); obtained by Monte Carlo simulation"],
       ["Individual player contribution", "Dropped: causal question, too little data; replaced by an ablation of data categories"]],
      widths=[6.5, 9.5], align_right_from=9)

# ================================================================== 2
h1("2. Data")
table([["Source", "Content", "Use"],
       ["football-data.co.uk", "Results, shots, corners, cards, bookmaker odds; 5 leagues, 2010/11 to 2026", "Core data"],
       ["Transfermarkt dataset (CC0)", "656,000 player market valuations", "Squad value"],
       ["Self-computed", "Elo ratings, shot-based \"xG proxy\"", "Features"]],
      widths=[4.2, 8.3, 3.5], align_right_from=9)
para("Volume: 29,239 matches (593 for OL), 189 clubs. Training 2011–2020 (about 18,200 matches), validation 2021–22 (3,651), test 2023 to October 2026 (5,602).")
para("@@Sources rejected.@@ FBref advanced statistics were removed by the provider in January 2026. Understat forbids automated access in its robots.txt. "
     "ClubElo was unreachable. There is no free historical source for injuries or line-ups, so I could not test them.")
para("@@Terms of use.@@ football-data.co.uk allows private, non-commercial use only. This is a private academic study; the raw data are downloaded, "
     "kept out of the repository and not redistributed. The Champions League fixtures come from Wikipedia (scores and dates only).")

# ================================================================== 3
h1("3. Variables (features) and how I avoid leakage")
para("Each variable is computed for the home team, the away team and their difference, @@using only matches before the one predicted.@@")
table([["Family", "Construction"],
       ["Elo rating", "Strength score updated after every match (formula below), one scale per league"],
       ["Recent form", "Average over the last 3 / 5 / 10 matches of points, goals for and against, shots on target"],
       ["Venue form", "Last 5 home (or away) matches"],
       ["League table", "Points and goal difference per match, position before the match"],
       ["Schedule", "Rest days, matches in the previous 14 days"],
       ["Squad value", "ln(total squad value) minus the league average of that month (Transfermarkt)"],
       ["xG proxy", "In practice 0.287 x shots on target (see below)"]],
      widths=[3.5, 12.5], align_right_from=9)
para("@@Elo.@@ Before a match, the expected score of the home team is:")
formula("E = 1 / (1 + 10^(−(R_home − R_away + 65) / 400))", "65 = home advantage in Elo points")
para("After the match, with S = 1 for a win, 0.5 for a draw and 0 for a loss and a multiplier for the goal difference:")
formula("R_home′ = R_home + K · (1 + ln(1 + |goal difference|)) · (S − E),     K = 20",
        "The rating stored for a match is the one BEFORE this update.")
para("@@Squad value@@ is measured relative to the league so that market inflation does not look like strength:")
formula("sv_rel = ln(V_team) − average of ln(V) over the league that month")
para("@@Leakage control.@@ Every rolling variable is shifted by one match; Elo is stored before its update; squad values use valuations dated "
     "before the snapshot; bookmaker odds are never a variable. An automated test replaces every result after a cutoff date by random values, "
     "re-runs the whole pipeline and requires that all earlier variables are strictly identical. I checked that the test fails on two bugs I injected on purpose "
     "(a rolling window without the shift; Elo stored after the update).")
para("@@Two lessons.@@ (1) The \"xG proxy\" collapses to 0.287 x shots on target: it is not a real xG. (2) Variables counting how many matches a team has been seen "
     "for are just a calendar clock, so I removed them (cost: +0.0011 log loss on validation, not significant).")

# ================================================================== 4
h1("4. The model I chose: multinomial logistic regression")
para("For a match with input vector x (100 variables), the model computes one linear score per outcome k (home win, draw, away win) and turns the three "
     "scores into probabilities with the softmax function:")
formula("P(outcome = k | x) = exp(wₖ · x + bₖ) / Σⱼ exp(wⱼ · x + bⱼ)",
        "wₖ = weights of outcome k, bₖ = intercept (it contains the home advantage); j runs over the 3 outcomes")
para("@@Estimation.@@ The weights are found by maximising the log-likelihood of the training matches (the same idea as the Solver exercise in the course), "
     "with an L2 penalty that keeps the weights small:")
formula("minimise   −(1/N) Σᵢ ln P(yᵢ | xᵢ)   +   (1 / 2CN) ‖w‖²")
para("The penalty strength 1/C is chosen by @@time-ordered cross-validation@@ inside the training seasons (each fold trains on the past and scores the next block). "
     "Inputs are standardised and missing values replaced by the median, both fitted on the training data only. The model is re-fitted once a year.")
h2("Why this model and not a more complex one")
bullets(["@@As accurate as the alternatives.@@ Boosting, random forest, SVM and neural networks were tested: none is significantly better (section 7).",
         "@@Explainable.@@ Each forecast splits exactly into the contribution of Elo, squad value, form and so on, because the model is linear before the softmax.",
         "@@Honest.@@ The probabilities are well calibrated: a 70% forecast comes true about 70% of the time.",
         "@@Fast.@@ It lets me simulate 10,000 seasons per league in minutes."])
h2("Models compared")
table([["Family", "Models"],
       ["Baselines", "Class frequencies; league table; Elo (calibrated to 3 classes); independent Poisson goals"],
       ["Linear (chosen)", "Multinomial logistic regression with L2 penalty"],
       ["Trees", "Random forest; gradient boosting"],
       ["Kernel", "RBF support-vector machine (under-tuned)"],
       ["Neural", "MLP with 1 and 2-3 hidden layers (L2, few epochs, average of 5 seeds)"],
       ["Online", "Logistic regression whose weights are updated by gradient descent after each week of results"],
       ["Reference", "Bet365 prices turned into probabilities (margin removed, average margin 5.5%)"]],
      widths=[3.5, 12.5], align_right_from=9)
para("Only scikit-learn and numpy are used (what we saw in the course). Dropout and Batch Normalisation do not exist in scikit-learn, "
     "so I replaced them by L2, few epochs and seed averaging.")

# ================================================================== 5
h1("5. How I measure quality")
para("Accuracy is a weak summary here: a draw is almost never the most likely outcome, so a model can be useful while rarely naming a draw. "
     "The main metric is the @@log loss@@, which rewards honest probabilities and punishes confident mistakes:")
formula("log loss = −(1/N) Σᵢ ln pᵢ", "pᵢ = probability the model gave to the outcome that actually happened in match i")
para("Example: if the model gave 20% to what happened, that match costs ln(5) = 1.61; if it gave 60%, it costs 0.51. "
     "A model that always says (44.7%, 25.2%, 30.0%) scores 1.075; a perfect-knowledge model would score 0.")
para("I also report the @@Brier score@@ and the @@calibration error@@ (ECE):")
formula("Brier = (1/N) Σᵢ Σₖ ( pᵢₖ − oᵢₖ )²", "oᵢₖ = 1 if outcome k happened in match i, 0 otherwise")
formula("ECE = Σ over bins of (n_b / N) · | average forecast_b − observed frequency_b |")
para("@@Significance.@@ Differences between models are tested with a paired block bootstrap on the log loss of each match, resampling whole calendar weeks "
     "(matches of the same week are not independent), 4,000 draws, 95% intervals. With 5,600 test matches, a difference below about 0.003 cannot be detected.")

# ================================================================== 6
h1("6. Validation")
para("Football data are a time series: a random split would let the model learn from matches played after the ones it is tested on.")
bullets(["Training 2011–2020, validation 2021–22 (used to compare models during development), test 2023 to October 2026 (used once, hyper-parameters frozen).",
         "@@Walk-forward:@@ each test season S is predicted by models fitted on seasons 2011 to S-1, as it would run in real life.",
         "Hyper-parameters are chosen only on the training seasons."])

# ================================================================== 7
h1("7. Results")
h2("7.1 Final test, walk-forward 2023 to October 2026 (5,602 matches)")
rows = [["Model", "Accuracy", "Log loss [95% CI]", "Brier", "ECE"]]
for name, label in [("Class frequency (naive)", "Baseline (frequencies)"), ("Elo (LogReg)", "Elo"), ("Poisson goals", "Poisson goals"),
                    ("LogReg (final)", "Logistic regression (chosen)"), ("Gradient boosting", "Gradient boosting"),
                    ("MLP deep (5 seeds)", "Neural network (MLP)"), ("Ensemble LogReg+GB+MLP", "Ensemble of the three"),
                    ("Bookmaker Bet365 (reference)", "Bookmaker (Bet365)")]:
    r = final.loc[name]
    rows.append([label, pct(r["accuracy"], 1), f"{r['log_loss']:.4f} [{r['log_loss_ci_low']:.3f}, {r['log_loss_ci_high']:.3f}]",
                 f"{r['brier']:.4f}", f"{r['ece']:.3f}"])
table(rows, widths=[5.2, 2.2, 5.2, 1.7, 1.7], bold_rows=(4,))
figure("final_logloss.png", 13.0, "Log loss on the test seasons: three model families are tied, the bookmaker is better.")
para("@@Reading.@@ The chosen logistic regression beats Elo (difference -0.0095, 95% CI [-0.0135, -0.0056]). Boosting (-0.0007) and the neural network (+0.0003) "
     "do @@not@@ beat it: their intervals contain 0. The bookmaker beats everyone (+0.0135 for my model, CI [+0.0094, +0.0175]). "
     "Complexity did not pay.")
h2("7.2 Calibration")
figure("final_reliability.png", 14.5, "Reliability diagrams: forecasts close to the diagonal are honest.")
para("The probabilities are well calibrated (ECE between 1.6% and 2.2%). The model almost never names a draw as the most likely outcome (6 matches out of 5,602): "
     "draws are essentially unpredictable.")
h2("7.3 Ablation: which data add value?")
table([["Model (logistic regression, validation)", "Log loss", "Change vs previous [95% CI]"],
       ["A: Elo only", "0.9957", "-"],
       ["B: + results form", "0.9931", "-0.0026 [-0.0051, -0.0001]"],
       ["C: + shots and xG proxy", "0.9887", "-0.0044 [-0.0072, -0.0017]"],
       ["D: + league table, fatigue", "0.9890", "+0.0003 [-0.0013, +0.0019]  (nothing)"],
       ["E: + squad value (final)", "0.9870", "-0.0020 [-0.0048, +0.0009]"]],
      widths=[6.6, 2.4, 7.0])
para("Form and shots each add a small, detectable gain on top of Elo; league position and fatigue add nothing. Squad value is not significant on two validation seasons, "
     "although it ranks second in the importance analysis: it overlaps with Elo (both measure long-run strength).")

# ================================================================== 8
h1("8. What drives the predictions?")
figure("phase7_family_importance.png", 12.0, "Increase in log loss when a family of variables is shuffled (validation seasons).")
para("@@Elo (+0.050) and squad value (+0.024) dominate.@@ Recent form adds +0.007, shots +0.003; rest days, league table and venue form are about 0. "
     "Home sides win 44.7% of matches, draw 25.2% and lose 30.0%, which sits in the intercept.")
para("@@Caution.@@ In the linear model, form has a slightly negative effect at fixed Elo and squad value. This is a collinearity artefact (overlapping windows), "
     "not evidence that a good run hurts. Importance describes what the model uses, not what causes results.")

# ================================================================== 9
h1("9. Model versus bookmaker")
para("An academic question only (no betting strategy is built): do my probabilities contain information that the market lacks? "
     "I mix the two by log-linear pooling and choose the weight w on validation:")
formula("p  ∝  p_market^(1 − w) · p_model^w")
para("The best weight on the model, measured on the test seasons, is @@w = 0.00@@: the market already contains everything the model knows. "
     "The pooled log loss (0.9663) is not better than the market alone (0.9658). The market takes a margin of 5.5% on average and is still sharper than my model: "
     "the likely reason is information I do not have (line-ups, injuries, news).")

# ================================================================== 10
h1("10. Previous results: was the model right?")
para("@@Important.@@ These forecasts are recomputed after the fact with the walk-forward method (model fitted only on earlier seasons, variables built only from earlier matches). "
     "They were @@not published in advance.@@ \"Right\" means the most likely outcome happened. A 45% pick is wrong more often than right, by design; "
     "log loss and calibration are the fair scores.")
rows = [["Period", "Matches", "Correct outcome", "Log loss: model", "Bookmaker", "No-skill"]]
for k, v in prev["summary"].items():
    rows.append([k, v["n"], pct(v["acc"]), f"{v['ll']:.3f}", f"{v['ll_book']:.3f}", f"{v['ll_naive']:.3f}"])
table(rows, widths=[5.2, 1.8, 2.5, 2.3, 2.4, 1.8])
h2("Do the percentages mean what they say?")
rows = [["Top pick given", "Matches", "Average forecast", "Came true"]]
for b in prev["calibration"]:
    rows.append([f"{int(b['lo'] * 100)}–{int(b['hi'] * 100)}%", b["n"], pct(b["said"]), pct(b["came_true"])])
table(rows, widths=[4.0, 3.0, 4.0, 3.0])
para("When the model gave 60–70% to its top pick, it happened 63% of the time. Over the last 30 days the model is right 56% of the time, a bit less sharp than the bookmaker.")
h2("Most recent matches")
names = ["Away win", "Draw", "Home win"]
rows = [["Date", "Match", "Score", "Home / draw / away", "Top pick", "Result"]]
for d, lg, h, a, hg, ag, y, p in prev["rows"][:14]:
    best = max(range(3), key=lambda i: p[i])
    rows.append([d[5:], f"{h} - {a}", f"{hg}-{ag}", f"{pct(p[2])} / {pct(p[1])} / {pct(p[0])}", f"{names[best]} {pct(p[best])}",
                 "Right" if best == y else "Wrong"])
table(rows, widths=[1.6, 5.0, 1.4, 3.5, 3.2, 1.4], size=9)

# ================================================================== 11
h1("11. OL case study")
figure("ol_strength_profile.png", 12.5, "OL's profile in Ligue 1 on 26 October 2026.")
para("OL is @@3rd in Ligue 1 by Elo@@ (1,615; Paris SG 1,707), 5th for recent points and goals scored, and 5th by squad value (EUR 305 m; defence only 10th). "
     "Results and rating are better than the shot volume (10th for shots on target) would suggest.")
para("@@Match prediction, Paris SG - Lyon (26 October 2026):@@")
table([["", "OL win", "Draw", "OL loss", "OL expected points"],
       ["Model", "15%", "21%", "64%", "0.67"],
       ["Bookmaker (margin removed)", "15%", "20%", "65%", "-"]], widths=[5.5, 2.2, 2.2, 2.2, 3.9])
para("Expected points are computed as:")
formula("E[points] = 3 · P(win) + 1 · P(draw)")
para("The 90% interval on OL's win probability is 14–17% with model uncertainty only (100 bootstrap refits) and 12–21% once the uncertainty on the true strength "
     "of both teams is added (Elo shaken by 60 points).")

# ================================================================== 12
h1("12. Season simulation (10,000 seasons per league)")
para("Each simulated season shakes the Elo of every team by a random shock, draws every remaining match outcome from the model's probabilities, draws a scoreline "
     "(for goal-difference tie-breaks) and ranks the table. The shock size (60 Elo points) was chosen by a @@backtest@@: each past season was simulated from 31 October "
     "using only what was known then, and compared with the real table.")
formula("Elo_team (in a simulation) = Elo_team + ε,     ε ~ Normal(0, 60²)")
table([["Backtest 2023–2025 (15 league-seasons)", "Simulator", "Naive reference"],
       ["Mean absolute error of final points", "6.4", "8.9"],
       ["Brier score, title", "0.022", "0.049"],
       ["Brier score, top 4", "0.064", "0.165"],
       ["Brier score, relegation", "0.055", "0.117"],
       ["Coverage of the 80% interval", "72%", "80%"]], widths=[8.0, 4.0, 4.0])
para("The simulator beats the naive reference, but its \"80%\" intervals cover only 72%: read them as lower bounds.")
h2("Projected top 3 of each league (expected final position)")
rows = [["League", "1st", "2nd", "3rd"]]
for L, name in LEAGUES.items():
    s = D["leagues"][L]["sim"]["summary"]
    top = sorted(s.items(), key=lambda kv: kv[1][3])[:3]
    rows.append([name] + [f"{t} ({v[0]:.0f} pts, title {pct(v[4])})" for t, v in top])
table(rows, widths=[3.0, 4.4, 4.4, 4.4], size=9)
h2("Ligue 1 2026/27 (state of 26 October, 7 matches played)")
rows = [["Team", "Expected points [80%]", "Title", "Champions League", "Top 4", "Relegation"]]
for t in ["Paris SG", "Monaco", "Lille", "Lyon"]:
    v = D["leagues"]["F1"]["sim"]["summary"][t]
    rows.append([t, f"{v[0]:.1f} [{v[1]:.0f}–{v[2]:.0f}]", pct(v[4], 1), pct(v[5], 1), pct(v[6], 1), pct(v[9], 1)])
table(rows, widths=[2.8, 3.8, 2.0, 3.2, 2.0, 2.2], bold_rows=(4,))
figure("sim_ol_position.png", 12.5, "OL's final position over 10,000 simulated seasons.")
para(f"OL's average final position is @@{ol[3]:.1f}@@, the most likely finishes are 2nd and 3rd ({pct(ol_pos[1])} and {pct(ol_pos[2])}), but the distribution is wide: "
     "after 7 matches six clubs share 13–14 points, so the forecast rests on team strength rather than on the table.")

# ================================================================== 13
h1("13. Champions League 2026/27")
para("Goal: use the real fixtures and combine each club's performance in its championship. The calendar (144 league-phase matches, 36 clubs) comes from Wikipedia; "
     f"only {euro_meta['played']} matches were recorded when I collected it, the others are simulated.")
para("@@European strength model.@@ Domestic leagues are only comparable through matches between their clubs, so I fitted a logistic regression on 3,085 past "
     "Champions League, Europa League and Conference League matches (2012–2025) with two differences between the clubs:")
formula("d_value = ln(squad value_home) − ln(squad value_away)")
formula("d_ppg = domestic points per match_home − domestic points per match_away", "domestic points per match = average over the last 20 league matches")
rows = [["Model (walk-forward 2023–2025, 943 matches)", "Log loss"]]
for k in ["Class frequency", "Squad value only", "Domestic form only", "Value + domestic form"]:
    rows.append([k, f"{euro.loc[k, 'log_loss']:.4f}"])
table(rows, widths=[10.0, 3.0])
para("@@Squad value carries the signal@@ (-0.051 vs frequencies, significant). @@Domestic form adds nothing measurable@@ once value is known (-0.0005, CI [-0.0048, +0.0037]); "
     "I keep it because the question asked to combine both, but its weight is small.")
para("@@Simulation.@@ 20,000 runs of the rest of the league phase, the play-offs (ranks 9–24), round of 16, quarter-finals, semi-finals and final.")
rows = [["Club", "Top 8", "Semi-final", "Final", "Winner"]]
for r in ucl.head(10).itertuples():
    rows.append([r.team, pct(r.top8), pct(r.SF), pct(r.Final), pct(r.Winner, 1)])
table(rows, widths=[5.5, 2.4, 2.8, 2.4, 2.4])
para("@@No clear favourite:@@ Arsenal, Manchester City, Barcelona, Real Madrid and Paris SG are each around 13%, Bayern 10%, Liverpool 6%. The differences between the first "
     "five are smaller than the uncertainty. This model is validated match by match, not on whole tournaments; the knockout bracket is approximate; AEK, Sabah and Slovan "
     "are not in the data and are assumed weak. No Europa League winner is predicted yet.")

# ================================================================== 14
h1("14. Limitations")
bullets(["No injuries, suspensions, line-ups or coaching news (the likely reason the bookmaker is better); no real xG.",
         "Squad values are a crowd-sourced snapshot (June 2026); 13 clubs have no value.",
         "Bookmaker benchmark: Bet365 pre-match prices with simple margin removal.",
         "Few validation windows; SVM under-tuned; online updating tested for the linear model only.",
         "Simulation: intervals too narrow in the backtest; UEFA places are assumptions; strength is constant apart from a random shock.",
         "Team state excludes results of 23–25 October 2026 that were not yet in the data source.",
         "football-data.co.uk limits use to private, non-commercial purposes."])

# ================================================================== 15
h1("15. Conclusion")
para("On public data, football results can be forecast with @@well-calibrated probabilities that are only modestly sharper than team strength alone@@. "
     "Elo and squad value explain most of what is predictable. A regularised logistic regression, my chosen model, is as good as boosting and neural networks: complexity did not pay. "
     "The models cannot beat the betting market and add no information to it, which is consistent with an efficient market. "
     f"For OL, the model sees a top-five squad with results slightly ahead of its shot volume: about {ol[0]:.1f} points, an average finish of {ol[3]:.1f}, "
     "and a coin-flip chance of the top four. So the disclaimer at the beginning stands: I really hope the ML is wrong.")

h1("Appendix: reproducibility")
para("Code, tests and the interactive website are in the GitHub repository (README). make pipeline rebuilds data, features, OL predictions and simulations; make all reruns every experiment; "
     "pytest runs 22 tests including the leakage test. All numbers come from files in reports/tables/.", size=10.5)


def main() -> None:
    doc.save(OUT)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(REPORTS), str(OUT)],
                   check=True, capture_output=True)
    print(OUT, OUT.with_suffix(".pdf"))


if __name__ == "__main__":
    main()
